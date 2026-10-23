from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from app.core.config import get_settings
from app.ingestion.common import ExtractionError, clean_text


class ExcelCell(BaseModel):
    sheet: str
    cell: str
    value: str
    number_format: str = ""


class ExtractedExcel(BaseModel):
    schema_version: int = 1
    cells: list[ExcelCell]


def extract_excel(path: Path) -> ExtractedExcel:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".xls":
        return extract_xls(path)
    raise ExtractionError("unsupported_type")


def extract_xlsx(path: Path) -> ExtractedExcel:
    max_cells = get_settings().extraction_max_cells
    with path.open("rb") as handle:
        workbook = load_workbook(handle, data_only=False, read_only=False)
    cells: list[ExcelCell] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    extracted = _xlsx_cell(cell, sheet.title)
                    if extracted is None:
                        continue
                    cells.append(extracted)
                    if len(cells) > max_cells:
                        raise ExtractionError("cell_limit_exceeded")
    finally:
        workbook.close()
    return ExtractedExcel(cells=cells)


def extract_xls(path: Path) -> ExtractedExcel:
    try:
        import xlrd
    except ImportError as error:
        raise ExtractionError("xls_unavailable") from error

    max_cells = get_settings().extraction_max_cells
    book = xlrd.open_workbook(path)
    cells: list[ExcelCell] = []
    for sheet in book.sheets():
        for row_index in range(sheet.nrows):
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                value = clean_text(str(cell.value))
                if value == "":
                    continue
                cells.append(
                    ExcelCell(
                        sheet=sheet.name,
                        cell=f"{get_column_letter(col_index + 1)}{row_index + 1}",
                        value=value,
                        number_format="",
                    )
                )
                if len(cells) > max_cells:
                    raise ExtractionError("cell_limit_exceeded")
    return ExtractedExcel(cells=cells)


def _xlsx_cell(cell: Cell, sheet: str) -> ExcelCell | None:
    if cell.value is None:
        return None
    value = clean_text(str(cell.value))
    if value == "":
        return None
    return ExcelCell(
        sheet=sheet,
        cell=f"{get_column_letter(cell.column)}{cell.row}",
        value=value,
        number_format=clean_text(str(cell.number_format or "")),
    )
