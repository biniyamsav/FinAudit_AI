from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from app.core.config import get_settings
from app.ingestion.common import ExtractionError
from app.ingestion.csv_text import extract_csv
from app.ingestion.docx_text import extract_docx
from app.ingestion.excel import extract_xls, extract_xlsx
from app.ingestion.pdf_text import extract_pdf_text
from tests.test_worker import make_pdf


def test_pdf_marks_empty_text_layer_as_needs_ocr(tmp_path: Path):
    path = tmp_path / "empty.pdf"
    path.write_bytes(make_pdf(""))
    extracted = extract_pdf_text(path)
    assert extracted.schema_version == 4
    assert extracted.pages[0].needs_ocr is True
    assert extracted.pages[0].page == 1


def test_pdf_page_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(get_settings(), "extraction_max_pages", 1)
    path = tmp_path / "two.pdf"
    path.write_bytes(_two_page_pdf())
    with pytest.raises(ExtractionError, match="page_limit_exceeded"):
        extract_pdf_text(path)


def test_xlsx_keeps_numbers_as_strings_with_format(tmp_path: Path):
    path = tmp_path / "book.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "P&L"
    sheet["B7"] = "1,234.50"
    sheet["C3"] = "Revenue"
    empty = book.create_sheet("Empty")
    empty["A1"] = None
    book.save(path)
    extracted = extract_xlsx(path)
    by_ref = {(cell.sheet, cell.cell): cell for cell in extracted.cells}
    assert extracted.schema_version == 1
    assert ("P&L", "B7") in by_ref
    assert by_ref[("P&L", "B7")].value == "1,234.50"
    assert isinstance(by_ref[("P&L", "B7")].value, str)
    assert by_ref[("P&L", "C3")].value == "Revenue"
    assert by_ref[("P&L", "B7")].number_format
    assert ("Empty", "A1") not in by_ref


def test_xlsx_cell_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(get_settings(), "extraction_max_cells", 1)
    path = tmp_path / "limit.xlsx"
    book = Workbook()
    book.active["A1"] = "a"
    book.active["A2"] = "b"
    book.save(path)
    with pytest.raises(ExtractionError, match="cell_limit_exceeded"):
        extract_xlsx(path)


def test_csv_sniffs_delimiter_and_strips_bom(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_bytes("Revenue;1000\nCost;250".encode("utf-8-sig"))
    extracted = extract_csv(path)
    assert extracted.schema_version == 1
    assert extracted.rows[0].row == 1
    assert extracted.rows[0].values == ["Revenue", "1000"]
    assert extracted.rows[1].values == ["Cost", "250"]
    assert all(isinstance(value, str) for row in extracted.rows for value in row.values)


def test_csv_row_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(get_settings(), "extraction_max_rows", 1)
    path = tmp_path / "rows.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ExtractionError, match="row_limit_exceeded"):
        extract_csv(path)


def test_docx_uses_paragraph_index_and_tables(tmp_path: Path):
    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Opening")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "1000"
    table.cell(1, 0).text = "Tax"
    table.cell(1, 1).text = "150"
    document.add_paragraph("Closing")
    document.save(path)
    extracted = extract_docx(path)
    texts = [item.text for item in extracted.paragraphs]
    assert "Opening" in texts
    assert "Closing" in texts
    assert extracted.paragraphs[0].paragraph_index == 1
    assert extracted.tables[0].rows[0] == ["Item", "1000"]
    assert extracted.tables[0].rows[1][1] == "150"
    assert isinstance(extracted.tables[0].rows[1][1], str)


def test_xls_extracts_cells(tmp_path: Path):
    try:
        import xlwt
    except ImportError:
        pytest.skip("xlwt is not installed, so a .xls fixture cannot be generated")

    path = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Summary")
    sheet.write(0, 1, "Revenue 1500")  # B1
    book.save(str(path))

    extracted = extract_xls(path)
    by_ref = {(cell.sheet, cell.cell): cell for cell in extracted.cells}
    assert extracted.schema_version == 1
    assert ("Summary", "B1") in by_ref
    assert by_ref[("Summary", "B1")].value == "Revenue 1500"


def _two_page_pdf() -> bytes:
    contents = [
        b"BT /F1 12 Tf 72 720 Td (Page one) Tj ET",
        b"BT /F1 12 Tf 72 720 Td (Page two) Tj ET",
    ]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 5 0 R /Resources << /Font << /F1 7 0 R >> >> >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 6 0 R /Resources << /Font << /F1 7 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(contents[0]) + contents[0] + b"\nendstream",
        b"<< /Length %d >>\nstream\n" % len(contents[1]) + contents[1] + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objects) + 1,
        xref_at,
    )
    return out
