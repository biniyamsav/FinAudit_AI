import uuid
from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import get_settings
from app.db.models import Document, ExtractedContent, IngestionJob
from app.db.session import get_session_factory
from app.main import app
from app.workers.ingestion_worker import process_document

client = TestClient(app)


def make_pdf(text: str) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objects) + 1,
        xref_at,
    )
    return out


def test_worker_extracts_pdf_text_and_deletes_temp_file():
    workspace_id = client.post("/api/v1/workspaces").json()["id"]
    files = {"file": ("r.pdf", make_pdf("Revenue 1000"), "application/pdf")}
    upload = client.post(f"/api/v1/workspaces/{workspace_id}/documents", files=files)
    document_id = uuid.UUID(upload.json()["id"])
    temp_file = Path(get_settings().temp_upload_dir) / workspace_id / str(document_id)
   
  

    with get_session_factory()() as db:
        document = db.get(Document, document_id)
        job = db.scalars(
            select(IngestionJob).where(IngestionJob.document_id == document_id)
        ).one()
        content = db.scalars(
            select(ExtractedContent).where(ExtractedContent.document_id == document_id)
        ).one()
        assert document.status == "ready"
        assert job.status == "succeeded"
        assert job.attempt_count == 1
        assert "Revenue 1000" in content.content["pages"][0]["text"]
        assert content.content["pages"][0]["needs_ocr"] is False
    assert not temp_file.exists()


def _xlsx_bytes() -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    book = Workbook()
    book.active["A1"] = "1000"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Note 1000")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def test_worker_extracts_xlsx():
    workspace_id = client.post("/api/v1/workspaces").json()["id"]
    files = {
        "file": (
            "book.xlsx",
            _xlsx_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    upload = client.post(f"/api/v1/workspaces/{workspace_id}/documents", files=files)
    document_id = uuid.UUID(upload.json()["id"])
    with get_session_factory()() as db:
        content = db.scalars(
            select(ExtractedContent).where(ExtractedContent.document_id == document_id)
        ).one()
        cells = content.content["cells"]
        assert cells[0]["cell"] == "A1"
        assert cells[0]["value"] == "1000"
        assert "number_format" in cells[0]


def test_worker_extracts_csv():
    workspace_id = client.post("/api/v1/workspaces").json()["id"]
    files = {"file": ("data.csv", b"Revenue,1000\n", "text/csv")}
    upload = client.post(f"/api/v1/workspaces/{workspace_id}/documents", files=files)
    document_id = uuid.UUID(upload.json()["id"])
    with get_session_factory()() as db:
        content = db.scalars(
            select(ExtractedContent).where(ExtractedContent.document_id == document_id)
        ).one()
        assert content.content["rows"][0]["values"] == ["Revenue", "1000"]


def test_worker_extracts_docx():
    workspace_id = client.post("/api/v1/workspaces").json()["id"]
    files = {
        "file": (
            "note.docx",
            _docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    upload = client.post(f"/api/v1/workspaces/{workspace_id}/documents", files=files)
    document_id = uuid.UUID(upload.json()["id"])
    with get_session_factory()() as db:
        content = db.scalars(
            select(ExtractedContent).where(ExtractedContent.document_id == document_id)
        ).one()
        assert content.content["paragraphs"][0]["text"] == "Note 1000"
        assert "paragraph_index" in content.content["paragraphs"][0]


def test_worker_keeps_temp_file_when_ocr_is_required():
    workspace_id = client.post("/api/v1/workspaces").json()["id"]
    files = {"file": ("scan.png", PNG_1X1, "image/png")}
    upload = client.post(f"/api/v1/workspaces/{workspace_id}/documents", files=files)
    document_id = uuid.UUID(upload.json()["id"])
    temp_file = Path(get_settings().temp_upload_dir) / workspace_id / str(document_id)
    with get_session_factory()() as db:
        document = db.get(Document, document_id)
        job = db.scalars(
            select(IngestionJob).where(IngestionJob.document_id == document_id)
        ).one()
        assert document.status == "failed"
        assert job.status == "failed"
        assert job.error_code == "no_text_found"
    assert temp_file.exists()